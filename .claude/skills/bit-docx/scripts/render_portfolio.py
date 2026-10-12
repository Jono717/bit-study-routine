#!/usr/bin/env python3
"""Render the BIT Portfolio of Evidence (or any Workings markdown draft) to .docx.

Two modes:

  * default (``--apa``, implied): the submission render. A4, 2.54 cm margins, Times New Roman 12,
    double-spaced, 0.5 inch first-line indents, APA 7 heading levels, page number top right,
    real Word TOC and Table of Figures fields, APA figure captions, APA tables, hanging-indent
    reference lists, and a "Back to Contents" link in the footer.
  * ``--review``: a reading render for the learner. Calibri 11, 1.15 spacing, landscape.

Formatting is the renderer's job; content is the learner's. The renderer never rewrites prose. It
does three structural things a consolidated portfolio needs, and reports each on stderr:

  1. Figures are renumbered consecutively across the whole portfolio, in order of first mention
     within each task, and in-text callouts ("Figure 2", "LT3 Figure 1") are rewritten to match.
     The markdown keeps its per-task numbers. Disable with ``--no-renumber``.
  2. The "Table of Contents" and "Table of Figures" sections are replaced by Word fields that
     Word fills on open (the document asks to update fields).
  3. Assembly artefacts are dropped: YAML front matter (used for document properties), the
     italic running-header line ("*Portfolio of Evidence ...*"), horizontal rules, and HTML
     comments.

Placeholders such as ``[NEEDS YOUR INPUT ...]``, ``[TODO ...]`` and ``[to confirm]`` are kept and
highlighted yellow. ``--submission`` refuses to render while any remain.

Landscape pages: wrap content in ``<!-- landscape -->`` ... ``<!-- /landscape -->`` comments. Tables
too tall for a portrait page (a wide matrix) move to landscape pages on their own, with the headings
that lead into them; ``--no-auto-landscape`` turns that off.

Tables are sized from the font's real glyph widths: every column is at least as wide as its longest
word, single-line columns get exactly their width, and wrapping columns share the rest by how much text
they hold.

Citations: ``[@key]`` and ``@key`` are formatted by Zotero's own citation engine from the Zotero library
export (see citations.py). Every "References" section is then generated: the consolidated list under the
level-1 References heading, and each task's own list under its level-2 heading (left out with
``--submission``, which keeps only the consolidated list). ``--bib FILE`` picks
the CSL JSON file; ``--no-citations`` leaves keys as typed.

Usage:
    uv run --with python-docx --with pillow python render_portfolio.py IN.md OUT.docx
        [--review] [--letter] [--toc-levels 1-2] [--no-renumber] [--no-footer-link] [--submission]
        [--no-auto-landscape] [--bib FILE] [--no-citations]
"""

from __future__ import annotations

import datetime
import pathlib
import re
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

import citations

LINK = re.compile(r"(?<!!)\[([^\]]+)\]\(([^)\s]+)\)")
INLINE = re.compile(r"(\*\*.+?\*\*|~~.+?~~|`[^`]+`|\*[^*\n]+?\*)")
PLACEHOLDER = re.compile(
    r"\[(?:NEEDS YOUR INPUT|TODO|to confirm|To be completed|to insert|TBD)[^\]]*\]", re.IGNORECASE
)
HEADING = re.compile(r"^(#{1,6})\s+(.*)")
ANCHOR = re.compile(r"\s*\{#([\w-]+)\}\s*$")
BULLET = re.compile(r"^\s*[-*]\s+(.*)")
NUMBERED = re.compile(r"^\s*(\d+)\.\s+(.*)")
SEPARATOR_ROW = re.compile(r"\|[\s\-:|]+\|")
IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$")
CAPTION = re.compile(r"\*\*(Figure|Table)\s+([A-Z]?\d+)\*\*")
APPENDIX = re.compile(r"Appendix\s+([A-Z])\b")
MENTION = re.compile(r"\b(Figure|Table)(s?)(\s+)(\d+)(?:(\s*(?:to|and|–|-)\s*)(\d+))?")
ITALIC_LINE = re.compile(r"\*[^*].*\*")
RUNNING_HEADER = re.compile(r"\*Portfolio of Evidence\b[^*]*\*")
BACK_LINK = re.compile(r"\[[^\]]*Back to Contents[^\]]*\]\(#[\w-]+\)", re.IGNORECASE)
LANDSCAPE_ON = re.compile(r"^\s*<!--\s*landscape\s*-->\s*$", re.IGNORECASE)
LANDSCAPE_OFF = re.compile(r"^\s*<!--\s*/landscape\s*-->\s*$", re.IGNORECASE)
FLATTEN_HEADERS = {("column", "content")}
TOC_HEADINGS = {"table of contents", "contents"}
TOF_HEADINGS = {"table of figures", "list of figures", "table of figures and tables", "list of figures and tables",
                "list of tables"}
REFERENCE_HEADINGS = {"references", "reference list"}
RENDERER = "bit-docx render_portfolio.py"
BREAK_MIN_WORDS = 120  # skip a page break if the current page holds less than this

FONT = {"apa": "Times New Roman", "review": "Calibri"}
FONT_FILES = {  # regular, bold; used only to measure table text
    "Times New Roman": ("/System/Library/Fonts/Supplemental/Times New Roman.ttf",
                        "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf"),
    "Calibri": ("/Applications/Microsoft Word.app/Contents/Resources/DFonts/Calibri.ttf",
                "/Applications/Microsoft Word.app/Contents/Resources/DFonts/Calibrib.ttf"),
}
CELL_SIZE = {"apa": 10, "review": 9}
CELL_PAD = Inches(0.08)  # left and right cell margin
TALL_ROW = Inches(2.5)   # a row taller than this on a portrait page sends its table to landscape
KEEP_ROW = Inches(2.0)   # rows shorter than this are never split across a page break
SHORT_CELL = 12          # a column whose body cells are all this short is centred
TBLPR_ORDER = ("tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
               "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook")
TCPR_ORDER = ("cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
              "textDirection", "tcFitText", "vAlign", "hideMark")
STYLES = {
    "body": "BIT Body",
    "title": "BIT Title",
    "cover": "BIT Cover Line",
    "caption": "BIT Figure Caption",
    "table_caption": "BIT Table Caption",
    "note": "BIT Note",
    "quote": "BIT Block Quote",
    "ref": "BIT Reference",
    "cell": "BIT Table Text",
    "back": "BIT Back Link",
    "figure": "BIT Figure",
}


def log(msg):
    print(msg, file=sys.stderr)


def slug(text):
    s = re.sub(r"[^\w\s-]", "", text.lower()).strip()
    return re.sub(r"\s+", "-", s)


def bookmark_name(anchor):
    return ("bm_" + re.sub(r"\W", "_", anchor))[:40]


# ---------------------------------------------------------------- parsing


def split_front_matter(lines):
    meta = {}
    if lines and lines[0].strip() == "---":
        for i, line in enumerate(lines[1:], 1):
            if line.strip() == "---":
                for kv in lines[1:i]:
                    if ":" in kv:
                        k, v = kv.split(":", 1)
                        meta[k.strip()] = v.strip().strip('"')
                return meta, lines[i + 1:]
    return meta, lines


def split_cover(lines):
    """If the document opens with an H1 title, return (title, cover lines, rest)."""
    i = 0
    while i < len(lines) and not lines[i].strip():
        i += 1
    if i >= len(lines) or not lines[i].startswith("# "):
        return None, [], lines
    title = lines[i][2:].strip()
    cover, j = [], i + 1
    while j < len(lines) and not lines[j].startswith("#"):
        s = lines[j].strip()
        if s and s != "---":
            cover.append(s)
        j += 1
    return title, cover, lines[j:]


def blocks(lines):
    """Group hard-wrapped markdown into blocks: (kind, payload)."""
    out, current = [], None

    def flush():
        nonlocal current
        if current:
            out.append(tuple(current))
            current = None

    in_comment = False
    for raw in lines:
        line = raw.rstrip()
        s = line.strip()
        if in_comment:
            in_comment = "-->" not in s
            continue
        if LANDSCAPE_ON.match(s) or LANDSCAPE_OFF.match(s):
            flush()
            out.append(("landscape", bool(LANDSCAPE_ON.match(s))))
            continue
        if s.startswith("<!--"):
            flush()
            in_comment = "-->" not in s
            continue
        if not s:
            flush()
            continue
        if m := IMAGE.match(s):
            flush()
            out.append(("image", (m.group(1), m.group(2))))
            continue
        if s.startswith("|"):
            if out and out[-1][0] == "table" and current is None:
                out[-1][1].append(s)
            else:
                flush()
                out.append(("table", [s]))
            continue
        if s == "---":
            flush()
            continue
        if m := HEADING.match(s):
            flush()
            text = m.group(2)
            am = ANCHOR.search(text)
            anchor = am.group(1) if am else None
            text = ANCHOR.sub("", text).strip()
            out.append(("heading", (len(m.group(1)), text, anchor or slug(text))))
            continue
        if s.startswith(">"):
            text = s.lstrip(">").strip()
            if current and current[0] == "quote":
                current[1] = f"{current[1]} {text}".strip()
            else:
                flush()
                current = ["quote", text]
            continue
        if m := BULLET.match(line):
            flush()
            current = ["bullet", m.group(1)]
            continue
        if m := NUMBERED.match(line):
            flush()
            current = ["numbered", m.group(2)]
            continue
        if current is None or current[0] == "quote":
            flush()
            current = ["para", s]
        else:
            current[1] = f"{current[1]} {s}"
    flush()
    return out


def drop_artefacts(items):
    kept, dropped = [], 0
    for kind, payload in items:
        if kind == "para" and RUNNING_HEADER.fullmatch(payload.strip()):
            dropped += 1
            continue
        kept.append((kind, payload))
    if dropped:
        log(f"dropped {dropped} running-header line(s)")
    return kept


def renumber_figures(items):
    """Renumber captions and callouts consecutively, by first mention within each task.

    Figures and tables under an "Appendix X" heading are numbered X1, X2 (APA 7) and stay out of the
    main sequence, so the figure list reads in page order.
    """
    sections, cur = [], []
    for item in items:
        if item[0] == "heading" and item[1][0] == 1 and cur:
            sections.append(cur)
            cur = []
        cur.append(item)
    sections.append(cur)

    def texts(item):
        kind, payload = item
        if kind in ("para", "bullet", "numbered", "quote"):
            return [payload]
        if kind == "heading":
            return [payload[1]]
        if kind == "table":
            return payload
        return []

    def task_of(section):
        for kind, payload in section:
            if kind == "heading" and payload[0] == 1:
                m = re.match(r"(?:Learning Task|LT)\s*(\d)", payload[1])
                if m:
                    return int(m.group(1))
        return None

    mapping, by_task, counter, report = {}, {}, {"Figure": 0, "Table": 0}, []
    for idx, section in enumerate(sections):
        task = task_of(section)
        for label in ("Figure", "Table"):
            captioned, letter_of, appendix = [], {}, None
            for kind, payload in section:
                if kind == "heading":
                    if appendix and payload[0] <= appendix[1]:
                        appendix = None
                    if m := APPENDIX.match(payload[1]):
                        appendix = (m.group(1), payload[0])
                elif (kind == "para" and (m := CAPTION.fullmatch(payload.strip())) and m.group(1) == label
                      and m.group(2).isdigit()):
                    captioned.append(int(m.group(2)))
                    if appendix:
                        letter_of[int(m.group(2))] = appendix[0]
            if not captioned:
                continue
            order = []
            for item in section:
                if item[0] == "para" and CAPTION.fullmatch(item[1].strip()):
                    continue
                for t in texts(item):
                    for m in re.finditer(rf"(?<!LT\d ){label}\s+(\d+)", t):
                        n = int(m.group(1))
                        if n in captioned and n not in order:
                            order.append(n)
            order += [n for n in captioned if n not in order]
            local, per_letter = {}, {}
            for n in order:
                if n in letter_of:
                    letter = letter_of[n]
                    per_letter[letter] = per_letter.get(letter, 0) + 1
                    local[n] = f"{letter}{per_letter[letter]}"
                else:
                    counter[label] += 1
                    local[n] = counter[label]
            mapping[(idx, label)] = local
            if task:
                by_task[(task, label)] = local
            report.append(f"{label}s in {'LT' + str(task) if task else 'section ' + str(idx)}: "
                          + ", ".join(f"{a}->{b}" for a, b in local.items()))

    unresolved = []

    def fix(text, idx):
        """Rewrite "Figure N", "Figures N to M" and "LT3 Figure N" to the portfolio-wide numbers.

        Which task's numbering applies: an "LTn" immediately before the label; else this section's own
        captions; else the last "LTn" earlier in the same sentence or table cell ("LT3 Appendix A and
        Figure 1"). A mention none of these resolves is left alone and reported, because after
        renumbering it may point at a different figure.
        """
        def repl(m):
            label, plural, space, a, sep, b = m.groups()
            before = text[:m.start()]
            if cross := re.search(r"LT(\d)\s+$", before):
                table = by_task.get((int(cross.group(1)), label))
            elif (idx, label) in mapping:
                table = mapping[(idx, label)]
            else:
                clause = re.split(r"(?<=[.!?])\s+|\|", before)[-1]
                tasks = re.findall(r"\bLT(\d)\b", clause)
                table = by_task.get((int(tasks[-1]), label)) if tasks else None
                if table is None and any(k[1] == label for k in by_task):
                    unresolved.append(f"...{before[-40:].lstrip()}{m.group(0)}")
            if not table:
                return m.group(0)
            if not b:
                return f"{label}{plural}{space}{table.get(int(a), int(a))}"
            span = range(int(a), int(b) + 1) if sep.strip() != "and" else (int(a), int(b))
            vals = [table.get(n, n) for n in span]
            if all(isinstance(v, int) for v in vals) and (sep.strip() == "and" or vals == list(range(vals[0], vals[-1] + 1))):
                return f"{label}{plural}{space}{vals[0]}{sep}{vals[-1]}"
            return f"{label}s{space}" + ", ".join(map(str, vals[:-1])) + f" and {vals[-1]}"

        return MENTION.sub(repl, text)

    out, skip = [], None  # the manual contents and figure lists are replaced by fields: leave them
    for idx, section in enumerate(sections):
        for kind, payload in section:
            if kind == "heading":
                if skip is not None and payload[0] <= skip:
                    skip = None
                if payload[1].lower().strip() in TOC_HEADINGS | TOF_HEADINGS:
                    skip = payload[0]
            if skip is not None:
                out.append((kind, payload))
                continue
            if kind in ("para", "bullet", "numbered", "quote"):
                payload = fix(payload, idx)
            elif kind == "heading":
                payload = (payload[0], fix(payload[1], idx), payload[2])
            elif kind == "table":
                payload = [fix(r, idx) for r in payload]
            out.append((kind, payload))
    for line in report:
        log("renumbered " + line)
    for ref in dict.fromkeys(unresolved):
        log(f"unresolved reference, may now point at the wrong figure (write 'LT3 Figure 1'): {ref}")
    return out


# ---------------------------------------------------------------- styling


def force_font(rpr_owner, name, size=None, lang="en-NZ"):
    """Set all four font slots and drop theme fonts, which otherwise override the name."""
    rpr = rpr_owner.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if lang:
        el = rpr.find(qn("w:lang"))
        if el is None:
            el = OxmlElement("w:lang")
            rpr.append(el)
        el.set(qn("w:val"), lang)
    if size:
        for tag in ("w:sz", "w:szCs"):
            el = rpr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                rpr.append(el)
            el.set(qn("w:val"), str(int(size * 2)))


def doc_defaults_rpr(doc):
    styles = doc.styles.element
    dd = styles.find(qn("w:docDefaults"))
    rpr_default = dd.find(qn("w:rPrDefault"))

    class Owner:
        def get_or_add_rPr(self):
            rpr = rpr_default.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                rpr_default.append(rpr)
            return rpr

    return Owner()


def para_style(doc, name, base="Normal"):
    try:
        return doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles[base]
        st.quick_style = True
        return st


def setup_styles(doc, mode):
    font = FONT[mode]
    apa = mode == "apa"
    size = 12 if apa else 11
    spacing = 2 if apa else 1.15

    force_font(doc_defaults_rpr(doc), font, size)
    normal = doc.styles["Normal"]
    force_font(normal.element, font, size)
    pf = normal.paragraph_format
    pf.line_spacing = spacing
    pf.space_before = Pt(0)
    pf.space_after = Pt(0 if apa else 6)
    pf.widow_control = True

    body = para_style(doc, STYLES["body"])
    body.paragraph_format.first_line_indent = Inches(0.5) if apa else None

    # APA 7: L1 centred bold; L2 flush-left bold; L3 flush-left bold italic;
    # L4 indented bold; L5 indented bold italic. All the body font and size, black.
    levels = [
        (WD_ALIGN_PARAGRAPH.CENTER, False, 0),
        (WD_ALIGN_PARAGRAPH.LEFT, False, 0),
        (WD_ALIGN_PARAGRAPH.LEFT, True, 0),
        (WD_ALIGN_PARAGRAPH.LEFT, False, 0.5),
        (WD_ALIGN_PARAGRAPH.LEFT, True, 0.5),
    ]
    for level, (align, italic, indent) in enumerate(levels, 1):
        st = doc.styles[f"Heading {level}"]
        force_font(st.element, font, size if apa else (14 if level == 1 else 12 if level == 2 else 11))
        st.font.bold = True
        st.font.italic = italic
        st.font.color.rgb = RGBColor(0, 0, 0)
        f = st.paragraph_format
        f.alignment = align
        f.left_indent = Inches(indent) if indent else None
        f.first_line_indent = None
        f.line_spacing = spacing
        f.space_before = Pt(0 if apa else 10)
        f.space_after = Pt(0 if apa else 4)
        f.keep_with_next = True

    title = para_style(doc, STYLES["title"])
    title.font.bold = True
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(96 if apa else 48)
    cover = para_style(doc, STYLES["cover"])
    cover.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    caption = para_style(doc, STYLES["caption"])
    caption.paragraph_format.keep_with_next = True
    caption.paragraph_format.space_before = Pt(6 if apa else 8)
    figure = para_style(doc, STYLES["figure"])
    figure.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.line_spacing = 1
    figure.paragraph_format.keep_with_next = True
    table_caption = para_style(doc, STYLES["table_caption"], base=STYLES["caption"])
    table_caption.paragraph_format.keep_with_next = True
    note = para_style(doc, STYLES["note"])
    note.paragraph_format.space_after = Pt(6 if apa else 8)

    quote = para_style(doc, STYLES["quote"])
    quote.paragraph_format.left_indent = Inches(0.5)
    if not apa:
        quote.font.italic = True

    ref = para_style(doc, STYLES["ref"])
    ref.paragraph_format.left_indent = Inches(0.5)
    ref.paragraph_format.first_line_indent = Inches(-0.5)

    cell = para_style(doc, STYLES["cell"])
    cell.font.size = Pt(CELL_SIZE[mode])
    cell.paragraph_format.line_spacing = 1
    cell.paragraph_format.space_after = Pt(2)
    cell.paragraph_format.first_line_indent = None

    back = para_style(doc, STYLES["back"])
    back.font.size = Pt(10)
    back.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    back.paragraph_format.line_spacing = 1

    for name in ("List Bullet", "List Number"):
        doc.styles[name].paragraph_format.line_spacing = spacing


# ---------------------------------------------------------------- inline runs


def add_hyperlink(paragraph, text, url=None, anchor=None):
    h = OxmlElement("w:hyperlink")
    if url:
        h.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    else:
        h.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    h.append(r)
    paragraph._p.append(h)


def leaf(paragraph, text, flags, size):
    pieces, last = [], 0
    for m in PLACEHOLDER.finditer(text):
        pieces.append((text[last:m.start()], False))
        pieces.append((m.group(0), True))
        last = m.end()
    pieces.append((text[last:], False))
    for piece, hot in pieces:
        if not piece:
            continue
        run = paragraph.add_run(piece)
        run.bold = flags.get("bold") or None
        run.italic = flags.get("italic") or None
        if flags.get("strike"):
            run.font.strike = True
        if flags.get("code"):
            run.font.name = "Menlo"
            run.font.size = Pt(10)
        if size:
            run.font.size = Pt(size)
        if hot:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def add_runs(paragraph, text, size=None, anchors=None, **flags):
    """Write markdown inline text: links, **bold**, *italic*, `code`, ~~strike~~, placeholders."""
    pos = 0
    for m in LINK.finditer(text):
        add_emphasis(paragraph, text[pos:m.start()], size, flags)
        label, target = m.group(1), m.group(2)
        if target.startswith("#"):
            if anchors is not None and target[1:] in anchors:
                add_hyperlink(paragraph, label, anchor=bookmark_name(target[1:]))
            else:
                add_emphasis(paragraph, label, size, flags)
        else:
            add_hyperlink(paragraph, label, url=target)
        pos = m.end()
    add_emphasis(paragraph, text[pos:], size, flags)


def add_emphasis(paragraph, text, size, flags):
    if not text:
        return
    parts = INLINE.split(text)
    if len(parts) == 1:
        leaf(paragraph, text, flags, size)
        return
    for token in parts:
        if not token:
            continue
        inner, nested = token, dict(flags)
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            inner, nested["bold"] = token[2:-2], True
        elif token.startswith("~~") and token.endswith("~~"):
            inner, nested["strike"] = token[2:-2], True
        elif token.startswith("`") and token.endswith("`"):
            inner, nested["code"] = token[1:-1], True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            inner, nested["italic"] = token[1:-1], True
        if inner == token:
            leaf(paragraph, token, nested, size)
        else:
            add_emphasis(paragraph, inner, size, nested)


# ---------------------------------------------------------------- fields, bookmarks, lists


def add_field(paragraph, instr, placeholder=""):
    def fld(kind):
        r = paragraph.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        if kind == "begin":
            el.set(qn("w:dirty"), "true")
        r._r.append(el)

    fld("begin")
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r._r.append(it)
    fld("separate")
    if placeholder:
        paragraph.add_run(placeholder).italic = True
    fld("end")


def bookmark(paragraph, name, ident):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(ident))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(ident))
    paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
    paragraph._p.append(end)


class Lists:
    """Restart List Number at 1 for every separate numbered list."""

    def __init__(self, doc):
        self.doc = doc
        try:
            num_pr = doc.styles["List Number"].element.pPr.numPr
            self.base = num_pr.numId.val
            self.numbering = doc.part.numbering_part.element
            nums = self.numbering.findall(qn("w:num"))
            base = next(n for n in nums if n.get(qn("w:numId")) == str(self.base))
            self.abstract = base.find(qn("w:abstractNumId")).get(qn("w:val"))
            self.next_id = max(int(n.get(qn("w:numId"))) for n in nums) + 1
        except Exception:  # template without numbering: fall back to the style's own list
            self.base = None
        self.current = None

    def start(self):
        if self.base is None:
            return
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(self.next_id))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), self.abstract)
        num.append(a)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), "1")
        ov.append(so)
        num.append(ov)
        self.numbering.append(num)
        self.current = self.next_id
        self.next_id += 1

    def apply(self, paragraph):
        if self.current is None:
            return
        num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = self.current


# ---------------------------------------------------------------- tables and images


def usable_width(doc):
    s = doc.sections[-1]
    return s.page_width - s.left_margin - s.right_margin


class Measure:
    """Width of a string in EMU in the table font, from the font file when it is installed."""

    def __init__(self, font, size):
        self.size, self.faces = size, None
        try:
            from PIL import ImageFont

            self.faces = tuple(ImageFont.truetype(f, size * 20) for f in FONT_FILES[font])
        except Exception:  # no Pillow or no font file: an average glyph is about half an em
            log(f"table widths estimated: no {font} font file to measure with")

    def __call__(self, text, bold=False):
        if self.faces:
            return int(self.faces[bold].getlength(text) * 12700 / 20)
        return int(len(text) * self.size * (0.55 if bold else 0.5) * 12700)

    @property
    def line(self):
        return int(self.size * 1.15 * 12700)


_measures = {}


def measure_for(mode):
    if mode not in _measures:
        _measures[mode] = Measure(FONT[mode], CELL_SIZE[mode])
    return _measures[mode]


def plain(text):
    return re.sub(r"\*\*|\*|`|~~", "", LINK.sub(r"\1", text)).strip()


def table_rows(raw_rows):
    rows = [[c.strip() for c in r.strip().strip("|").split("|")]
            for r in raw_rows if not SEPARATOR_ROW.fullmatch(r.strip())]
    ncols = max((len(r) for r in rows), default=0)
    return [r + [""] * (ncols - len(r)) for r in rows]


def flattened(rows):
    return len(rows[0]) == 2 and tuple(c.strip("*").lower() for c in rows[0]) in FLATTEN_HEADERS


def table_layout(rows, usable, measure, pad=2 * CELL_PAD + Pt(3)):
    """Column widths, centred columns and estimated lines per row for a table at a given text width.

    Every column is at least as wide as its longest word, so nothing breaks mid-word. Columns whose
    text fits on one line get exactly that width; the rest of the width goes to the wrapping columns
    in proportion to how much text each holds, which evens out the line count across a row.
    """
    n = len(rows[0])
    lo, hi, vol = [pad] * n, [pad] * n, [1] * n
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            t, bold = plain(text), r == 0
            lo[c] = max(lo[c], max((measure(w, bold) for w in t.split()), default=0) + pad)
            hi[c] = max(hi[c], measure(t, bold) + pad)
            if r:
                vol[c] += measure(t)
    if sum(hi) <= usable:
        widths = [h + (usable - sum(hi)) * h // sum(hi) for h in hi]
    elif sum(lo) >= usable:
        widths = [usable * w // sum(lo) for w in lo]
    else:
        widths, flex = list(lo), list(range(n))
        while flex:
            spare = usable - sum(widths[c] for c in range(n) if c not in flex) - sum(lo[c] for c in flex)
            total = sum(vol[c] for c in flex)
            share = {c: lo[c] + spare * vol[c] // total for c in flex}
            over = [c for c in flex if share[c] > hi[c]]
            for c in over:
                widths[c] = hi[c]
                flex.remove(c)
            if not over:
                for c in flex:
                    widths[c] = share[c]
                break
    centre = [c > 0 and all(len(plain(row[c])) <= SHORT_CELL for row in rows[1:]) for c in range(n)]
    return widths, centre, row_lines(rows, widths, measure, pad)


def row_lines(rows, widths, measure, pad=2 * CELL_PAD + Pt(3)):
    """Estimated line count of the tallest cell in each row (10% allowance for word wrap)."""
    return [max(1, max(-(-int(measure(plain(t), r == 0) * 1.1) // max(w - pad, 1))
                       for t, w in zip(row, widths)))
            for r, row in enumerate(rows)]


def header_key(rows):
    return tuple(plain(c).lower() for c in rows[0])


def put(parent, el, order):
    """Insert a property element where the schema expects it, replacing any existing one."""
    tag = el.tag.split("}")[1]
    for old in parent.findall(el.tag):
        parent.remove(old)
    rank = order.index(tag)
    for child in parent:
        name = child.tag.split("}")[1]
        if name in order and order.index(name) > rank:
            child.addprevious(el)
            return
    parent.append(el)


def border(parent, edge, val="single", sz=8, color="000000"):
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), val)
    if val != "nil":
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
    parent.append(el)


def apa_borders(table, row_rules=False):
    """APA 7 tables: rules above and below the table and under the header, nothing vertical.

    Tables whose rows wrap also get a hairline grey rule between rows. APA allows horizontal rules
    where they aid clarity, and without them long multi-line rows run into each other.
    """
    b = OxmlElement("w:tblBorders")
    for edge, val in (("top", "single"), ("left", "nil"), ("bottom", "single"), ("right", "nil")):
        border(b, edge, val)
    border(b, "insideH", "single" if row_rules else "nil", sz=4, color="A6A6A6")
    border(b, "insideV", "nil")
    put(table._tbl.tblPr, b, TBLPR_ORDER)
    for cell in table.rows[0].cells:
        tb = OxmlElement("w:tcBorders")
        border(tb, "bottom", "single", sz=6)
        put(cell._tc.get_or_add_tcPr(), tb, TCPR_ORDER)


def fix_geometry(table, widths, vpad):
    """Make the grid, the table width and the cell margins explicit and agree with the cell widths,
    so Word, Pages and LibreOffice lay the table out the same way."""
    tbl = table._tbl
    for col, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(round(w / 635)))
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(round(w / 635) for w in widths)))
    tbl_w.set(qn("w:type"), "dxa")
    put(tbl.tblPr, tbl_w, TBLPR_ORDER)
    mar = OxmlElement("w:tblCellMar")
    for edge, size in (("top", vpad), ("left", CELL_PAD), ("bottom", vpad), ("right", CELL_PAD)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(int(size) // 635))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    put(tbl.tblPr, mar, TBLPR_ORDER)


def row_flag(row, tag):
    el = OxmlElement(tag)
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def emit_table(doc, raw_rows, mode, anchors, family=None):
    """``family``: every row of every table sharing this header, so consecutive tables of one kind
    (for example a matrix split across several tables) get identical column widths."""
    rows = table_rows(raw_rows)
    if not rows:
        return
    if flattened(rows):
        for row in rows[1:]:
            label, content = row[:2]
            p = doc.add_paragraph(style=STYLES["body"] if mode == "apa" else None)
            add_runs(p, label.strip(), anchors=anchors, bold=True)
            add_runs(doc.add_paragraph(style=STYLES["body"]), content.strip(), anchors=anchors)
        return
    measure = measure_for(mode)
    widths, centre, _ = table_layout(family or rows, usable_width(doc), measure)
    lines = row_lines(rows, widths, measure)
    wraps = any(n > 1 for n in lines[1:])
    table = doc.add_table(rows=len(rows), cols=len(widths))
    table.style = "Table Grid"
    table.autofit = False
    fix_geometry(table, widths, Inches(0.04) if wraps else Inches(0.015))
    for r, row in enumerate(rows):
        tr = table.rows[r]
        if r == 0:
            row_flag(tr, "w:tblHeader")
        if lines[r] * measure.line < KEEP_ROW:
            row_flag(tr, "w:cantSplit")
        for c, text in enumerate(row):
            cell = tr.cells[c]
            cell.width = widths[c]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.BOTTOM if r == 0 else WD_ALIGN_VERTICAL.TOP
            p = cell.paragraphs[0]
            p.style = doc.styles[STYLES["cell"]]
            if centre[c]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if r == 0:
                p.paragraph_format.keep_with_next = True
            add_runs(p, text.strip(), anchors=anchors, bold=(r == 0) or None)
    if mode == "apa":
        apa_borders(table, row_rules=wraps)
    doc.add_paragraph(style=STYLES["note"])


def auto_landscape(items, portrait, measure):
    """Put tables too tall for a portrait page on landscape pages, with the headings that lead in.

    A wide self-assessment matrix is the case in point. Tables with the same header
    as a tall one go with it, and tall tables separated only by headings and a few lines share one
    landscape run, so the reader turns the page once. Content already inside landscape markers is
    left alone.
    """
    inside, candidates, tall_headers = False, [], set()
    for i, (kind, payload) in enumerate(items):
        if kind == "landscape":
            inside = payload
        elif kind == "table" and not inside:
            rows = table_rows(payload)
            if rows and not flattened(rows):
                header = header_key(rows)
                candidates.append((i, header))
                _, _, lines = table_layout(rows, portrait, measure)
                if max(lines[1:], default=0) * measure.line > TALL_ROW:
                    tall_headers.add(header)
    tall = [i for i, header in candidates if header in tall_headers]

    def leads_in(item):
        kind, payload = item
        if kind == "heading":
            return payload[0] > 1
        return kind == "para" and (CAPTION.fullmatch(payload.strip()) or
                                   (ITALIC_LINE.fullmatch(payload.strip()) and len(payload.split()) < 25))

    def bridge(gap):
        return (all((k == "heading" and p[0] > 1) or k == "para" for k, p in gap)
                and sum(len(p.split()) for k, p in gap if k == "para") <= 150)

    runs = []
    for t in tall:
        start = t
        while start > 0 and leads_in(items[start - 1]):
            start -= 1
        if runs and bridge(items[runs[-1][1] + 1:start]):
            runs[-1][1] = t
        else:
            runs.append([start, t])
    out = list(items)
    for start, end in reversed(runs):
        out.insert(end + 1, ("landscape", False))
        out.insert(start, ("landscape", True))
    if runs:
        log(f"landscape pages: {len(tall)} tall table(s) in {len(runs)} run(s)")
    return out


def rasterise_svg(svg, out, px=2000, pad=16):
    """Render SVG to PNG via QuickLook (no cairo or rsvg on this machine), trimming the padding."""
    from PIL import Image, ImageChops, ImageOps

    txt = pathlib.Path(svg).read_text(encoding="utf-8")
    w = float(re.search(r'width="([\d.]+)"', txt).group(1))
    h = float(re.search(r'height="([\d.]+)"', txt).group(1))
    side = max(w, h)
    padded = re.sub(r'(<svg[^>]*?)height="[\d.]+"', lambda m: m.group(1) + f'height="{side}"', txt, count=1)
    padded = re.sub(r'viewBox="[^"]+"', f'viewBox="0 0 {side} {side}"', padded, count=1)
    with tempfile.TemporaryDirectory() as td:
        tmp = pathlib.Path(td) / "padded.svg"
        tmp.write_text(padded, encoding="utf-8")
        subprocess.run(["qlmanage", "-t", "-s", str(px), "-o", td, str(tmp)], capture_output=True)
        rendered = next(pathlib.Path(td).glob("*.png"), None)
        if rendered is None:
            raise RuntimeError(f"QuickLook produced no thumbnail for {svg}")
        im = Image.open(rendered).convert("RGB")
    blank = Image.new("RGB", im.size, (255, 255, 255))
    diff = ImageChops.difference(im, blank).convert("L").point(lambda v: 255 if v > 8 else 0)
    im = ImageOps.expand(im.crop(diff.getbbox()), border=pad, fill=(255, 255, 255))
    im.save(out, optimize=True)


def emit_image(doc, alt, src, base_dir, cache, max_height=Inches(4.6)):
    from PIL import Image

    path = (base_dir / src).resolve()
    if not path.exists():
        add_runs(doc.add_paragraph(style=STYLES["note"]), f"[missing image: {src}]", italic=True)
        log(f"missing image: {src}")
        return
    if path.suffix.lower() == ".svg":
        cache.mkdir(parents=True, exist_ok=True)
        png = cache / f"{path.stem}.png"
        if not png.exists() or png.stat().st_mtime < path.stat().st_mtime:
            rasterise_svg(path, png)
        path = png
    width = usable_width(doc)
    with Image.open(path) as im:
        ratio = im.height / im.width
    if width * ratio > max_height:
        width = int(max_height / ratio)
    p = doc.add_paragraph(style=STYLES["figure"])
    shape = p.add_run().add_picture(str(path), width=int(width))
    shape._inline.docPr.set("descr", alt or path.stem)
    shape._inline.docPr.set("title", alt or path.stem)


# ---------------------------------------------------------------- page furniture


def page_setup(section, mode, letter):
    if letter:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    else:
        section.page_width, section.page_height = Mm(210), Mm(297)
    if mode == "review":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        margin = Inches(0.75)
    else:
        margin = Inches(1)
    section.left_margin = section.right_margin = margin
    section.top_margin = section.bottom_margin = margin


def page_number(section, mode):
    if mode == "apa":
        p = section.header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_field(p, "PAGE", "1")
    else:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("Page ")
        add_field(p, "PAGE", "1")
        p.add_run(" of ")
        add_field(p, "NUMPAGES", "1")


def footer_link(section, anchor):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_hyperlink(p, "Back to Contents", anchor=bookmark_name(anchor))
    for r in p._p.iter(qn("w:r")):
        rpr = r.find(qn("w:rPr"))
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "18")
        rpr.append(sz)


def set_landscape(doc, on):
    prev = doc.sections[-1]
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    want = WD_ORIENT.LANDSCAPE if on else WD_ORIENT.PORTRAIT
    if sec.orientation != want:
        sec.orientation = want
        sec.page_width, sec.page_height = prev.page_height, prev.page_width


# ---------------------------------------------------------------- main


def convert(src, dst, mode="apa", letter=False, toc_levels="1-2", renumber=True, footer=True,
            submission=False, landscape=True, cite=True, bib_path=None):
    src = pathlib.Path(src)
    lines = src.read_text(encoding="utf-8").splitlines()
    meta, lines = split_front_matter(lines)
    title, cover, lines = split_cover(lines)
    items = drop_artefacts(blocks(lines))
    if renumber:
        items = renumber_figures(items)
    if landscape and mode == "apa":
        portrait = (Inches(8.5) if letter else Mm(210)) - 2 * Inches(1)
        items = auto_landscape(items, portrait, measure_for(mode))
    bibliography, cited_by = None, {}
    if cite:
        library, is_seed = (pathlib.Path(bib_path), False) if bib_path else citations.find_library(src.parent)
        if library and (result := citations.apply(items, citations.load_library(library), log)):
            items, bibliography, cited_by = result
            log(f"library: {library}" + (" (the seed import file: export the Zotero collection to "
                                         "references/zotero_library.json to make Zotero the source)" if is_seed else ""))

    holes = sorted({m.group(0) for _, p in items if isinstance(p, str) for m in PLACEHOLDER.finditer(p)}
                   | {m.group(0) for line in cover for m in PLACEHOLDER.finditer(line)})
    if holes:
        log(f"{len(holes)} distinct placeholder(s) highlighted: " + "; ".join(h[:40] for h in holes[:8]))
        if submission:
            log("refusing to render a submission copy while placeholders remain")
            return 2

    doc = Document()
    setup_styles(doc, mode)
    lists = Lists(doc)
    page_setup(doc.sections[0], mode, letter)
    page_number(doc.sections[0], mode)
    doc.core_properties.title = meta.get("title", title or src.stem)
    doc.core_properties.author = meta.get("author", "")
    doc.core_properties.subject = "Bachelor of Information Technology, Portfolio of Evidence"
    # Provenance: Word replaces lastModifiedBy and adds w14:paraId on save, which is how the checker
    # tells an untouched render from a file the learner has edited in Word.
    now = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0)
    doc.core_properties.created = doc.core_properties.modified = now
    doc.core_properties.last_modified_by = RENDERER

    anchors = {p[2] for k, p in items if k == "heading"}
    has_table_captions = any(k == "para" and (m := CAPTION.fullmatch(p.strip())) and m.group(1) == "Table"
                             for k, p in items)
    families = {}
    for kind, payload in items:
        if kind == "table" and (rows := table_rows(payload)):
            families.setdefault(header_key(rows), [rows[0]]).extend(rows[1:])
    cache = src.parent / ".figure_cache"
    words_on_page = 0
    first_page = True

    if title:
        add_runs(doc.add_paragraph(style=STYLES["title"]), title)
        for line in cover:
            add_runs(doc.add_paragraph(style=STYLES["cover"]), line)
        words_on_page = BREAK_MIN_WORDS
        first_page = False

    skip_until_level = None
    ref_level = None
    generated_level = None  # a References section filled from Zotero: the typed entries under it are skipped
    section = None
    in_front = True
    list_kind = None
    pending_caption = None
    toc_anchor = None
    bm_id = 0

    def maybe_break(paragraph, force=False):
        nonlocal words_on_page
        if (force or words_on_page >= BREAK_MIN_WORDS) and not first_page:
            paragraph.paragraph_format.page_break_before = True
            words_on_page = 0

    for i, (kind, payload) in enumerate(items):
        if kind == "heading":
            level, text, anchor = payload
            if skip_until_level is not None and level <= skip_until_level:
                skip_until_level = None
            if ref_level is not None and level <= ref_level:
                ref_level = None
            if generated_level is not None and level <= generated_level:
                generated_level = None
            if level == 1:
                section = anchor
            if (bibliography and submission and level > 1 and text.lower().strip() in REFERENCE_HEADINGS
                    and skip_until_level is None):
                skip_until_level = level  # submission: one consolidated list after LT7, no per-task lists
        if skip_until_level is not None:
            continue
        if generated_level is not None and not (kind == "para" and PLACEHOLDER.search(payload)):
            continue
        if kind != "numbered":
            list_kind = None

        if kind == "landscape":
            set_landscape(doc, payload)
            words_on_page = 0
            continue

        if kind == "heading":
            level, text, anchor = payload
            key = text.lower().strip()
            p = doc.add_heading(level=min(level, 5))
            add_runs(p, text, anchors=anchors)
            bm_id += 1
            bookmark(p, bookmark_name(anchor), bm_id)
            if level == 1 and re.match(r"(learning task|lt)\s*\d", key):
                in_front = False
            if level == 1 and re.match(r"(learning task \d|references?\b|appendi)", key):
                maybe_break(p, force=True)
            elif level == 1 or (in_front and level == 2):
                maybe_break(p)
            if key in REFERENCE_HEADINGS:
                ref_level = level
                if bibliography:
                    keys = cited_by.get(section, set())
                    for ids, entry in bibliography:
                        if level == 1 or keys & set(ids):
                            add_runs(doc.add_paragraph(style=STYLES["ref"]), entry, anchors=anchors)
                    generated_level = level
            if key in TOC_HEADINGS:
                toc_anchor = anchor
                add_field(doc.add_paragraph(), f'TOC \\o "{toc_levels}" \\h \\z \\u',
                          "Right-click and choose Update Field to build the table of contents.")
                skip_until_level = level
            elif key in TOF_HEADINGS:
                for kind_, style, what in (("Figure", "caption", "figures"), ("Table", "table_caption", "tables")):
                    if kind_ == "Table" and not has_table_captions:
                        continue
                    if has_table_captions:
                        add_runs(doc.add_paragraph(style=STYLES["note"]), kind_ + "s", bold=True)
                    add_field(doc.add_paragraph(), f'TOC \\h \\z \\t "{STYLES[style]},1"',
                              f"Right-click and choose Update Field to build the list of {what}.")
                skip_until_level = level
            words_on_page += 5
            first_page = False
            continue

        first_page = False
        if kind == "table":
            rows = table_rows(payload)
            emit_table(doc, payload, mode, anchors, families.get(header_key(rows)) if rows else None)
            words_on_page += sum(len(r.split()) for r in payload)
        elif kind == "image":
            emit_image(doc, payload[0], payload[1], src.parent, cache)
            words_on_page += 150
        elif kind == "quote":
            add_runs(doc.add_paragraph(style=STYLES["quote"]), payload, anchors=anchors)
            words_on_page += len(payload.split())
        elif kind == "bullet":
            add_runs(doc.add_paragraph(style="List Bullet"), payload, anchors=anchors)
            words_on_page += len(payload.split())
        elif kind == "numbered":
            p = doc.add_paragraph(style="List Number")
            if list_kind != "numbered":
                lists.start()
                list_kind = "numbered"
            lists.apply(p)
            add_runs(p, payload, anchors=anchors)
            words_on_page += len(payload.split())
        else:  # paragraph
            text = payload.strip()
            words_on_page += len(text.split())
            if cap := CAPTION.fullmatch(text):
                nxt = items[i + 1] if i + 1 < len(items) else None
                p = doc.add_paragraph(style=STYLES["caption" if cap.group(1) == "Figure" else "table_caption"])
                add_runs(p, text, anchors=anchors)
                if nxt and nxt[0] == "para" and ITALIC_LINE.fullmatch(nxt[1].strip()):
                    p.add_run().add_break(WD_BREAK.LINE)
                    add_runs(p, nxt[1].strip(), anchors=anchors)
                    pending_caption = i + 1
                continue
            if pending_caption == i:
                pending_caption = None
                continue
            if BACK_LINK.fullmatch(text):
                add_runs(doc.add_paragraph(style=STYLES["back"]), text, anchors=anchors)
                continue
            if text.startswith("*Note.*") or text.startswith("*Note:"):
                style = STYLES["note"]
            elif ref_level is not None:
                style = STYLES["ref"]
            else:
                style = STYLES["body"]
            add_runs(doc.add_paragraph(style=style), text, anchors=anchors)

    if footer and mode == "apa" and toc_anchor:
        footer_link(doc.sections[0], toc_anchor)

    settings = doc.settings.element
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.append(upd)

    doc.save(dst)
    log(f"wrote {dst}")
    return 0


if __name__ == "__main__":
    argv = sys.argv[1:]
    flags = {a for a in argv if a.startswith("--")}
    args = [a for a in argv if not a.startswith("--")]
    toc = next((argv[k + 1] for k, a in enumerate(argv) if a == "--toc-levels"), "1-2")
    bib = next((argv[k + 1] for k, a in enumerate(argv) if a == "--bib"), None)
    args = [a for a in args if a not in (toc, bib)]
    if len(args) != 2:
        print(__doc__)
        sys.exit(1)
    sys.exit(convert(args[0], args[1], mode="review" if "--review" in flags else "apa",
                     letter="--letter" in flags, toc_levels=toc, renumber="--no-renumber" not in flags,
                     footer="--no-footer-link" not in flags, submission="--submission" in flags,
                     landscape="--no-auto-landscape" not in flags, cite="--no-citations" not in flags,
                     bib_path=bib))
