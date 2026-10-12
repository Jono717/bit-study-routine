#!/usr/bin/env python3
"""Audit a portfolio .docx against the BIT / APA 7 formatting spec (references/format-spec.md).

Prints one row per check: PASS, WARN or FAIL, with the evidence. Exit code 1 if any FAIL.

Usage:
    uv run --with python-docx python check_docx.py FILE.docx [--submission]

``--submission`` turns placeholders and draft notes into FAILs.
"""

from __future__ import annotations

import re
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PLACEHOLDER = re.compile(r"\[(?:NEEDS YOUR INPUT|TODO|to confirm|To be completed|to insert|TBD)[^\]]*\]", re.I)
rows = []


def report(status, check, evidence):
    rows.append((status, check, evidence))


def resolve(paragraph, getter):
    """Walk direct formatting then the style chain for a paragraph-format attribute."""
    value = getter(paragraph.paragraph_format)
    style = paragraph.style
    while value is None and style is not None:
        value = getter(style.paragraph_format)
        style = style.base_style
    return value


def theme_fonts(rpr):
    if rpr is None:
        return False
    rf = rpr.find(qn("w:rFonts"))
    return rf is not None and any(k.endswith("Theme") or k.endswith("theme") for k in rf.attrib)


def table_geometry(d):
    """Column widths against the text: grid agreement, margins, mid-word breaks, tall portrait rows."""
    from render_portfolio import TALL_ROW, Measure  # same measurement the renderer sizes columns with

    section_of, k = [], 0
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:tbl"):
            section_of.append(k)
        elif el.tag == qn("w:p") and el.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
            k += 1
    try:
        size = d.styles["BIT Table Text"].font.size.pt
    except (KeyError, AttributeError):
        size = 10
    measure = Measure(d.styles["Normal"].font.name or "Times New Roman", size)
    mismatch, too_wide, mid_word, tall = [], [], [], []
    for n, (t, sec_no) in enumerate(zip(d.tables, section_of), 1):
        tbl = t._tbl
        grid = [int(g.get(qn("w:w")) or 0) * 635 for g in tbl.tblGrid.findall(qn("w:gridCol"))]
        try:
            cells = [[c for c in row.cells] for row in t.rows]
        except Exception:
            continue
        if any(len(row) != len(grid) for row in cells):
            continue  # merged cells: geometry checks do not apply
        first = [c.width for c in cells[0]]
        if any(w is None or abs(w - g) > 30 * 635 for w, g in zip(first, grid)):
            mismatch.append(f"T{n}")
        widths = [w or g for w, g in zip(first, grid)]
        sec = d.sections[min(sec_no, len(d.sections) - 1)]
        if sum(widths) > sec.page_width - sec.left_margin - sec.right_margin + Inches(0.2):
            too_wide.append(f"T{n}")
        mar = tbl.tblPr.find(qn("w:tblCellMar")) if tbl.tblPr is not None else None
        sides = [mar.find(qn(f"w:{s}")) if mar is not None else None for s in ("left", "right")]
        pad = sum(int(e.get(qn("w:w"))) * 635 if e is not None else 108 * 635 for e in sides)  # Word default 0.075 in
        tallest = 0
        for r, row in enumerate(cells):
            row_lines = 1
            for c, cell in enumerate(row):
                text = cell.text.strip()
                need = max((measure(w, r == 0) for w in text.split()), default=0) + pad
                name = f"T{n} '{cells[0][c].text.strip()[:14]}'"
                if need > widths[c] + Pt(1) and name not in mid_word:
                    mid_word.append(name)
                row_lines = max(row_lines, -(-int(measure(text) * 1.1) // max(widths[c] - pad, 1)))
            tallest = max(tallest, row_lines * measure.line if r else 0)
        if sec.page_width < sec.page_height and tallest > TALL_ROW:
            tall.append(f"T{n} ({Emu(tallest).inches:.1f} in row)")
    report("PASS" if not mismatch else "WARN", "Table grid agrees with cell widths", ", ".join(mismatch) or "all")
    report("PASS" if not too_wide else "WARN", "Tables fit within the margins", ", ".join(too_wide) or "all")
    report("PASS" if not mid_word else "WARN", "No table column narrower than its longest word",
           "; ".join(mid_word[:5]) + (f" (+{len(mid_word) - 5})" if len(mid_word) > 5 else "") if mid_word else "none")
    report("PASS" if not tall else "WARN", "Tall tables on landscape pages", "; ".join(tall) or "none")


def main(path, submission=False):
    d = Document(path)
    paras = d.paragraphs
    body_xml = d.element.xml
    with zipfile.ZipFile(path) as z:
        parts = {n: z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.endswith(".xml")}
    headers = "".join(v for k, v in parts.items() if "header" in k or "footer" in k)

    by_word = "w14:paraId" in body_xml
    who = d.core_properties.last_modified_by or "unknown"
    report("INFO", "Provenance", "saved in Microsoft Word since rendering: do not overwrite without asking"
           if by_word else f"untouched render (last modified by {who})")

    # page
    s = d.sections[0]
    w, h = Emu(s.page_width).mm, Emu(s.page_height).mm
    size = "A4" if abs(min(w, h) - 210) < 2 else "Letter" if abs(min(w, h) - 215.9) < 2 else f"{w:.0f}x{h:.0f}mm"
    report("PASS" if size == "A4" else "WARN", "Paper size (NZ submission: A4)", f"{size}, {len(d.sections)} section(s)")
    margins = {round(Emu(getattr(s, m)).inches, 2) for m in ("left_margin", "right_margin", "top_margin", "bottom_margin")}
    report("PASS" if margins == {1.0} else "FAIL", "Margins 2.54 cm / 1 inch", f"{sorted(margins)} in")

    # fonts
    dd = d.styles.element.find(qn("w:docDefaults"))
    dd_rpr = dd.find(".//" + qn("w:rPr")) if dd is not None else None
    leaks = [n for n in ("Normal", "Heading 1", "Heading 2", "Heading 3")
             if theme_fonts(d.styles[n].element.rPr)]
    if theme_fonts(dd_rpr):
        leaks.insert(0, "docDefaults")
    report("FAIL" if leaks else "PASS", "No theme fonts overriding Times New Roman", ", ".join(leaks) or "none")
    normal = d.styles["Normal"]
    report("PASS" if normal.font.name == "Times New Roman" and normal.font.size and normal.font.size.pt == 12 else "FAIL",
           "Normal is Times New Roman 12", f"{normal.font.name} {normal.font.size.pt if normal.font.size else '?'} pt")
    report("PASS" if normal.paragraph_format.line_spacing == 2 else "FAIL", "Double spacing",
           f"Normal line spacing {normal.paragraph_format.line_spacing}")
    other = {r.font.name for p in paras for r in p.runs if r.font.name and r.font.name not in ("Times New Roman", "Menlo")}
    report("WARN" if other else "PASS", "No stray run-level fonts", ", ".join(sorted(other)) or "none")

    # headings
    heads = [(p, int(p.style.name.split()[-1])) for p in paras if p.style.name.startswith("Heading ")]
    by_level = {}
    for _, lvl in heads:
        by_level[lvl] = by_level.get(lvl, 0) + 1
    report("PASS" if heads else "FAIL", "Heading styles used", ", ".join(f"H{k}: {v}" for k, v in sorted(by_level.items())))
    h1 = [p for p, lvl in heads if lvl == 1]
    dupes, dupe_ids, idx = [], set(), {id(p._p): i for i, p in enumerate(paras)}
    def task_no(text):
        m = re.match(r"(?:Learning Task|LT)\s*(\d)", text)
        return m.group(1) if m else None

    for a, b in zip(h1, h1[1:]):
        between = paras[idx[id(a._p)] + 1: idx[id(b._p)]]
        same_task = task_no(a.text) and task_no(a.text) == task_no(b.text)
        if same_task and sum(len(x.text.split()) for x in between) < 120:  # renderer's BREAK_MIN_WORDS
            dupes.append(f'"{a.text[:30]}" / "{b.text[:30]}"')
            dupe_ids.add(id(b._p))
    expected = [p for p in h1 if id(p._p) not in dupe_ids]
    broken = sum(1 for p in expected if p.paragraph_format.page_break_before)
    report("PASS" if expected and broken == len(expected) else "WARN", "Level-1 headings start a new page",
           f"{broken} of {len(expected)} (duplicate openers excluded: {len(dupe_ids)})")
    report("WARN" if dupes else "PASS", "No back-to-back duplicate level-1 headings",
           "; ".join(dupes[:4]) + (f" (+{len(dupes) - 4})" if len(dupes) > 4 else "") if dupes else "none")
    blue = [n for n in ("Heading 1", "Heading 2", "Heading 3", "Heading 4")
            if d.styles[n].font.color and d.styles[n].font.color.type and str(d.styles[n].font.color.rgb) != "000000"]
    report("FAIL" if blue else "PASS", "Headings black (APA)", ", ".join(blue) or "all black")

    # body
    body = [p for p in paras if not p.style.name.startswith(("Heading", "List", "BIT Figure", "BIT Note", "BIT Reference",
                                                          "BIT Block", "BIT Title", "BIT Cover", "BIT Back", "BIT Table"))
            and len(p.text.split()) >= 20]
    indented = sum(1 for p in body if (resolve(p, lambda f: f.first_line_indent) or 0) > 0)
    ratio = indented / len(body) if body else 0
    report("PASS" if ratio >= 0.8 else "FAIL", "Body paragraphs have a 0.5 inch first-line indent",
           f"{indented} of {len(body)} ({ratio:.0%})")
    singles = sum(1 for p in body if len(re.findall(r"[.!?](\s|$)", p.text.strip())) <= 1)
    report("WARN" if singles else "PASS", "Single-sentence paragraphs", f"{singles}")

    # fields
    report("PASS" if "TOC \\o" in body_xml else "FAIL", "Word TOC field (auto-updating contents)",
           "present" if "TOC \\o" in body_xml else "missing; contents typed by hand")
    has_tof = bool(re.search(r"TOC \\[hz][^<]*\\[tc]", body_xml))
    report("PASS" if has_tof else "FAIL", "Table of Figures field", "present" if has_tof else "missing")
    upd = "updateFields" in parts.get("word/settings.xml", "")
    report("PASS" if upd else "WARN", "Word asked to update fields on open", "yes" if upd else "no")
    report("PASS" if " PAGE " in headers else "FAIL", "Page numbers", "PAGE field present" if " PAGE " in headers else "none")

    # figures
    captions = [p for p in paras if p.style.name.startswith("BIT Figure Caption") or
                (re.fullmatch(r"Figure \d+", p.text.strip().split("\n")[0]) and p.runs and p.runs[0].bold)]
    labels = [m.group(1) for p in captions if (m := re.match(r"Figure ([A-Z]?\d+)", p.text.strip()))]
    nums = [int(x) for x in labels if x.isdigit()]
    lettered = [x for x in labels if not x.isdigit()]
    status = "FAIL" if sorted(nums) != list(range(1, len(nums) + 1)) else "WARN" if nums != sorted(nums) else "PASS"
    report(status, "Figure numbers consecutive 1..N, in page order",
           f"{nums}" + (f"; appendix figures {lettered}" if lettered else "") + (" (out of page order)" if status == "WARN" else ""))
    merged = sum(1 for p in captions if "\n" in p.text or any(br.get(qn("w:type")) in (None, "textWrapping")
                                                              for br in p._p.iter(qn("w:br"))))
    report("PASS" if captions and merged == len(captions) else "WARN",
           "APA caption: bold number, title on the next line (soft return)", f"{merged} of {len(captions)}")
    tables_seen, captioned = 0, 0
    prev = None
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:tbl"):
            tables_seen += 1
            text = "".join(t.text or "" for t in prev.iter(qn("w:t"))) if prev is not None and prev.tag == qn("w:p") else ""
            style = prev.find(qn("w:pPr") + "/" + qn("w:pStyle")) if prev is not None and prev.tag == qn("w:p") else None
            sid = style.get(qn("w:val")) if style is not None else ""
            if "TableCaption" in sid or re.match(r"Table [A-Z]?\d+", text):
                captioned += 1
        prev = el
    report("PASS" if captioned == tables_seen else "WARN", "Tables carry an APA caption (Table N, italic title)",
           f"{captioned} of {tables_seen}")
    if captioned:
        has_lot = bool(re.search(r'TOC \\h \\z \\t "BIT Table Caption', body_xml))
        report("PASS" if has_lot else "WARN", "List of Tables field", "present" if has_lot else "missing")
    shapes = d.inline_shapes
    no_alt = sum(1 for sh in shapes if not sh._inline.docPr.get("descr"))
    report("PASS" if not no_alt else "FAIL", "Images carry alt text", f"{len(shapes) - no_alt} of {len(shapes)}")

    # tables
    wide_cells, vertical = 0, 0
    for t in d.tables:
        tbl_pr = t._tbl.tblPr
        b = tbl_pr.find(qn("w:tblBorders")) if tbl_pr is not None else None
        iv = b.find(qn("w:insideV")) if b is not None else None
        if iv is None or iv.get(qn("w:val")) not in ("nil", "none"):
            vertical += 1
        cell_p = t.cell(min(1, len(t.rows) - 1), 0).paragraphs[0]
        if (resolve(cell_p, lambda f: f.line_spacing) or 1) >= 1.5:
            wide_cells += 1
    report("PASS" if not wide_cells else "FAIL", "Table cells single-spaced", f"{wide_cells} of {len(d.tables)} tables double-spaced")
    report("PASS" if not vertical else "WARN", "APA table rules (no vertical lines)", f"{vertical} of {len(d.tables)} tables have vertical lines")
    table_geometry(d)

    # references
    in_ref, ref_paras, hung = None, 0, 0
    for p in paras:
        if p.style.name.startswith("Heading "):
            lvl = int(p.style.name.split()[-1])
            if p.text.strip().lower() in ("references", "reference list"):
                in_ref = lvl
                continue
            if in_ref is not None and lvl <= in_ref:
                in_ref = None
        if (in_ref is not None and len(p.text.split()) >= 5 and "Back to Contents" not in p.text
                and not p.style.name.startswith("List")):
            ref_paras += 1
            fli = resolve(p, lambda f: f.first_line_indent)
            if fli is not None and fli < 0:
                hung += 1
    report("PASS" if ref_paras and hung == ref_paras else "FAIL", "Reference entries have a 0.5 inch hanging indent",
           f"{hung} of {ref_paras}")

    # lists and quotes
    numbered = [p for p in paras if p.style.name == "List Number"]
    ids = {p._p.pPr.numPr.numId.val for p in numbered if p._p.pPr is not None and p._p.pPr.numPr is not None}
    report("PASS" if not numbered or ids else "WARN", "Numbered lists restart at 1",
           f"{len(numbered)} items, {len(ids) or 'one shared'} numbering instance(s)")
    blue_quote = sum(1 for p in paras if p.style.name == "Intense Quote")
    report("PASS" if not blue_quote else "WARN", "Block quotes use APA indentation, not Intense Quote", f"{blue_quote} Intense Quote")

    # leaks and content flags
    text = "\n".join(p.text for p in paras) + "\n" + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    leaks = {k: len(re.findall(v, text, re.M)) for k, v in
             {"](": r"\]\(", "{#": r"\{#", "**": r"\*\*", "<!--": r"<!--", "YAML": r"^title:\s", "~~": r"~~",
              "[@key] (unformatted citation)": r"\[-?@\w"}.items()}
    leaks = {k: v for k, v in leaks.items() if v}
    report("FAIL" if leaks else "PASS", "No raw markdown in the document", ", ".join(f"{k} x{v}" for k, v in leaks.items()) or "none")
    holes = len(PLACEHOLDER.findall(text)) + len(re.findall(r"\bTODO\b", text))
    report(("FAIL" if submission else "WARN") if holes else "PASS", "No placeholders or TODOs", f"{holes}")
    em = text.count("—")
    report("WARN" if em else "PASS", "Em dashes (learner's own style rule)", f"{em}")

    width = max(len(c) for _, c, _ in rows)
    for status, check, evidence in rows:
        print(f"{status:<4}  {check:<{width}}  {evidence}")
    count = {k: sum(1 for r in rows if r[0] == k) for k in ("PASS", "WARN", "FAIL")}
    fails = count["FAIL"]
    print(f"\n{sum(count.values())} checks: {count['PASS']} pass, {count['WARN']} warn, {fails} fail")
    return 1 if fails else 0


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if len(args) != 1:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(args[0], submission="--submission" in sys.argv))
